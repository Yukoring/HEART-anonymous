"""
Build the recording workbook for the real-robot farm run.

Five sheets:

  Overview        what to run and in what order
  A_Grasp         the 60 independent grasp trials, one row each, ready to fill in
  B_Plan          all 45 plans, with the failure predicted offline
  B_Plan_Actions  every action of every plan, raw and as PDDL, to follow along
  Scenes          the placement tables, so the sheet is self-contained

All three iterations of each scene-condition pair are executed, so no plan has
to be chosen over another and the same pair's run-to-run variation is measured
rather than assumed.

    python -m experiments.build_farm_workbook
"""

import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from experiments.generate_farm_scenes import INSTRUCTION
from heart.evaluation.feasibility_oracle import get_capability, graspable

PROJECT_ROOT = Path(__file__).parent.parent
# Pinned to a run rather than "the newest one" — the sheet is printed and taken
# to the robot, so the plans in it have to stay the plans it was built from.
# This run is the one whose farm prompt states the action signatures only and
# leaves every physical judgement to the planner.
RUN_DIR = PROJECT_ROOT / "results" / "farm_planners_20260831_144651"
OUT = PROJECT_ROOT / "results" / "farm_experiment_sheet.xlsx"
ROBOT = "summit_ur5e"

CONDITIONS = [("heart_llm_cot", "LLM-CoT + HEART"),
              ("baseline_llm_cot", "LLM-CoT alone"),
              ("baseline_triple_s", "Triple-S")]

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
BAND = PatternFill("solid", fgColor="EEF2F8")
FILLIN = PatternFill("solid", fgColor="FFF7CC")   # cells the operator writes
BAD = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    # Set by reference, not via ws.cell() — that call materialises the cell and
    # would push every later append() down by one row.
    ws.freeze_panes = f"A{row + 1}"


def widths(ws, spec: Dict[str, int]) -> None:
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def load_scenes() -> Dict[int, List[dict]]:
    scenes = {}
    for seed in range(1, 6):
        path = PROJECT_ROOT / "data" / "farm" / "scenes" / f"farm_seed{seed:02d}_scene_graph.json"
        scene = list(json.load(open(path)).values())[0]
        rows = []
        for stem in sorted(r for r in scene["rooms"] if r.startswith("stem")):
            for name, item in sorted(scene["rooms"][stem]["items"].items()):
                if "pick" not in item.get("affordance", []):
                    continue
                ok, why = graspable(ROBOT, item)
                rows.append({
                    "name": name, "stem": stem,
                    "z": round(item["location"][2], 3),
                    "w": round(min(item["size"]), 3),
                    "ripeness": item["state"][0],
                    "graspable": ok,
                    "violates": "" if ok else why[0].split("(")[0],
                })
        scenes[seed] = rows
    return scenes


def load_runs() -> List[dict]:
    runs = []
    for row in csv.DictReader(open(RUN_DIR / "picks.csv")):
        if row["error"]:
            continue
        seed, it = int(row["seed"]), row["iteration"]
        stem = RUN_DIR / "plans" / f"farm_seed{seed:02d}_iter{it}_{row['condition']}"
        raw = stem.with_name(stem.name + ".plan")
        pddl = stem.with_name(stem.name + "_pddl.plan")
        read = lambda p: [l.strip() for l in p.read_text().splitlines() if l.strip()] if p.is_file() else []
        plan = read(raw)
        runs.append({
            "seed": seed, "iteration": it, "condition": row["condition"],
            "plan": plan, "pddl": read(pddl),
            "picks": row["picks"].split("|") if row["picks"] else [],
            "infeasible": row["infeasible"].split("|") if row["infeasible"] else [],
            "valid": row["valid"] == "True",
            "info": " ".join(row["info"].split()),
        })
    return runs


def fail_step(run: dict) -> int:
    """
    1-based index of the action VAL rejected, or 0 if the plan ran to the end.

    VAL reports "at time N" for a sequential plan, and the converter emits one
    PDDL action per input action, so N indexes both plans alike. A goal that is
    simply unmet carries no time — nothing failed, the plan just stopped short.
    """
    if run["valid"]:
        return 0
    match = re.search(r"at time (\d+)", run["info"])
    return int(match.group(1)) if match else 0


def diagnose(run: dict, tomatoes: List[dict]) -> tuple:
    """(예상 결과, 원인 분류, 상세) — offline VAL verdict, read for an operator."""
    if run["valid"]:
        return "완주", "—", "오프라인 검증 통과"
    if not run["plan"]:
        return "실패", "플랜 없음", "플래너가 액션을 생성하지 못함"
    if run["infeasible"]:
        first = run["infeasible"][0]
        spec = next((t for t in tomatoes if t["name"] == first), None)
        why = {"too_high": "도달 불가", "too_wide": "너무 큼"}.get(spec["violates"], "물리 제약") if spec else "물리 제약"
        detail = f"{first} ({why}"
        if spec:
            detail += f", z={spec['z']} m, 최소치수={spec['w']} m)"
        else:
            detail += ")"
        return "실패", "물리 제약 위반", detail

    match = re.search(r"Set \(item_collected (\w+)\)", run["info"])
    if match:
        return "실패", "수집 누락", f"{match.group(1)}를 수집하지 않음"
    if "agent_loaded" in run["info"]:
        return "실패", "적재 누락", "손에 든 채로 다음 대상을 집으려 함"
    return "실패", "목표 미달성", run["info"][:70]


def main() -> int:
    scenes = load_scenes()
    runs = load_runs()
    cap = get_capability(ROBOT)
    ordered = sorted(runs, key=lambda r: (r["seed"],
                     [c for c, _ in CONDITIONS].index(r["condition"]), r["iteration"]))
    labels = dict(CONDITIONS)

    wb = Workbook()

    # ---------------------------------------------------------------- Overview
    ws = wb.active
    ws.title = "Overview"
    widths(ws, {"A": 22, "B": 96})
    lines = [
        ("농장 실기 실험 기록지", ""),
        ("", ""),
        ("로봇", f"Summit XL + UR5e — 도달 높이 {cap.reach_height:.3f} m, "
                 f"그리퍼 개폐 {cap.gripper_opening:.4f} m, 정격 페이로드 {cap.payload} kg"),
        ("지시문", INSTRUCTION),
        ("", ""),
        ("측정 (A)", "독립 파지 60회 — 시트 A_Grasp. 오라클 예측이 실제와 맞는지 보는 1차 측정이며, "
                     "플래너와 무관하게 토마토 하나씩 독립적으로 시도합니다."),
        ("측정 (B)", "플랜 실행 45회 — 시트 B_Plan. 씬 5개 × 조건 3개 × 3회. "
                     "세 회차를 모두 실행하므로 어느 플랜을 고를지 정할 필요가 없고, "
                     "같은 조건의 회차 간 편차도 측정됩니다."),
        ("", ""),
        ("진행 순서", "씬 1을 세팅 → A_Grasp에서 그 씬의 토마토를 전부 시도 → "
                      "B_Plan에서 그 씬의 세 조건을 실행 → 씬 2로 이동"),
        ("", ""),
        ("노란 칸", "실기에서 채워 넣는 칸입니다. 나머지는 오프라인에서 확정된 값이므로 수정하지 마세요."),
        ("실패 원인", "도달 불가 / 파지 실패 / 충돌 / 주행 실패 중 하나로 적습니다. "
                      "앞의 셋이 Reviewer 17이 지적한 kinematic · dynamics · collision에 각각 대응합니다."),
        ("예상 실패 액션", "오프라인 검증(VAL)이 거부한 액션의 번호이며, B_Plan_Actions의 '#' 열과 "
                           "같은 번호입니다. 그 뒤의 액션은 실행에 도달하지 못하므로 '미도달'로 표시했습니다. "
                           "한 액션이 물리 제약과 다른 조건을 동시에 어길 때는 물리 쪽을 원인으로 적습니다 "
                           "— 실기에서 실제로 관측되는 것이 그쪽이기 때문입니다."),
        ("목표 미달성", "액션은 전부 성립하는데 목표만 못 채운 경우로, 지목할 액션이 없어 '—'로 둡니다. "
                        "이 플랜은 끝까지 실행되며, 무엇을 빠뜨렸는지가 기록 대상입니다."),
        ("", ""),
        ("(A) 중단 규칙", "오라클이 '가능'이라 한 토마토가 2회 중 1회라도 실패하면 3회차를 추가합니다."),
        ("(B) 중단 규칙", "goto가 실패하면 위치를 복구할 수 없으므로 그 자리에서 중단하고 첫 실패 지점만 기록합니다. "
                          "액션 단위 비율은 (B)에서 계산하지 않습니다 — 그것은 (A)가 담당합니다."),
    ]
    for i, (k, v) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=bool(k), size=11 if i == 1 else 10)
        c = ws.cell(row=i, column=2, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name=FONT, bold=True, size=14)

    # ----------------------------------------------------------------- A_Grasp
    ws = wb.create_sheet("A_Grasp")
    head = ["씬", "토마토", "줄기", "높이 (m)", "최소치수 (m)", "익음",
            "오라클 판정", "위반 축", "시행", "결과 (성공/실패)", "실패 원인", "비고"]
    ws.append(head)
    style_header(ws, 1, len(head))
    widths(ws, {"A": 6, "B": 12, "C": 10, "D": 10, "E": 13, "F": 9, "G": 12,
                "H": 11, "I": 7, "J": 16, "K": 16, "L": 24})
    row = 2
    for seed in sorted(scenes):
        for t in scenes[seed]:
            for trial in (1, 2):
                ws.append([seed, t["name"], t["stem"][:7], t["z"], t["w"], t["ripeness"],
                           "가능" if t["graspable"] else "불가능",
                           {"too_high": "높이", "too_wide": "크기"}.get(t["violates"], "—"),
                           trial, "", "", ""])
                for c in range(1, len(head) + 1):
                    cell = ws.cell(row=row, column=c)
                    cell.font = Font(name=FONT, size=10)
                    cell.border = BOX
                    cell.alignment = Alignment(horizontal="center")
                    if seed % 2 == 0:
                        cell.fill = BAND
                if not t["graspable"]:
                    ws.cell(row=row, column=7).fill = BAD
                for c in (10, 11, 12):
                    ws.cell(row=row, column=c).fill = FILLIN
                row += 1
    example = row
    ws.append(["예시", "tomato_02", "stem_02", 0.97, 0.067, "ripe", "가능", "—", 1,
               "성공", "—", "기록 형식 예시 — 실제 시행 아님"])
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=example, column=c)
        cell.font = Font(name=FONT, size=10, italic=True, color="808080")
        cell.alignment = Alignment(horizontal="center")

    # ------------------------------------------------------------------ B_Plan
    ws = wb.create_sheet("B_Plan")
    head = ["씬", "조건", "회차", "액션 수", "집으려는 대상", "오프라인 예상",
            "예상 실패 원인", "예상 실패 액션 #", "상세", "실기 완주 (O/X)",
            "첫 실패 액션", "실패 원인", "비고"]
    ws.append(head)
    style_header(ws, 1, len(head))
    widths(ws, {"A": 6, "B": 18, "C": 7, "D": 9, "E": 34, "F": 13,
                "G": 16, "H": 15, "I": 40, "J": 15, "K": 20, "L": 16, "M": 22})
    row = 2
    for run in ordered:
        label = dict(CONDITIONS)[run["condition"]]
        outcome, cause, detail = diagnose(run, scenes[run["seed"]])
        stop = fail_step(run)
        # No step number when the plan ran to completion and only the goal was
        # unmet — there is no action to point at.
        ws.append([run["seed"], label, int(run["iteration"]), len(run["plan"]),
                   ", ".join(run["picks"]) or "(없음)", outcome, cause,
                   f"#{stop} / {len(run['plan'])}" if stop else "—", detail,
                   "", "", "", ""])
        for c in range(1, len(head) + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=(c in (5, 9)))
            if run["seed"] % 2 == 0:
                cell.fill = BAND
        if outcome != "완주":
            ws.cell(row=row, column=6).fill = BAD
            if stop:
                ws.cell(row=row, column=8).font = Font(name=FONT, size=10, bold=True,
                                                       color="9C0006")
        for c in (10, 11, 12, 13):
            ws.cell(row=row, column=c).fill = FILLIN
        row += 1

    # ---------------------------------------------------------- B_Plan_Actions
    ws = wb.create_sheet("B_Plan_Actions")
    head = ["씬", "조건", "회차", "#", "액션 (플래너 출력)", "PDDL 변환", "대상 판정",
            "오프라인 예상", "실행 (O/X)", "실패 원인"]
    ws.append(head)
    style_header(ws, 1, len(head))
    widths(ws, {"A": 6, "B": 18, "C": 7, "D": 5, "E": 42, "F": 46, "G": 16,
                "H": 16, "I": 13, "J": 18})
    row = 2
    pick_re = re.compile(r"\b(?:pick|pick_from)\s*\(\s*[^,)]+\s*,\s*([^,)]+)")
    for run in ordered:
        label = dict(CONDITIONS)[run["condition"]]
        spec = {t["name"]: t for t in scenes[run["seed"]]}
        actions = run["plan"] or ["(플랜 없음)"]
        # The converter is told to emit one PDDL action per input action, so the
        # two lists line up; pad rather than assume it when a run fell short.
        pddl = run["pddl"] + [""] * (len(actions) - len(run["pddl"]))
        stop = fail_step(run)
        for i, (action, converted) in enumerate(zip(actions, pddl), start=1):
            m = pick_re.search(action)
            verdict = ""
            if m:
                t = spec.get(m.group(1).strip().strip(")'\""))
                if t:
                    verdict = "가능" if t["graspable"] else (
                        "불가능 (도달)" if t["violates"] == "too_high" else "불가능 (크기)")
            # Everything after the rejected action is unreachable in execution,
            # so it is marked rather than left blank — a blank there would read
            # as "expected to succeed".
            expect = "통과" if not stop or i < stop else (
                f"실패 지점 (#{i})" if i == stop else "미도달")
            ws.append([run["seed"], label, int(run["iteration"]), i, action,
                       converted, verdict, expect, "", ""])
            for c in range(1, len(head) + 1):
                cell = ws.cell(row=row, column=c)
                cell.font = Font(name=FONT, size=10)
                cell.border = BOX
                if run["seed"] % 2 == 0:
                    cell.fill = BAND
            if verdict.startswith("불가능"):
                ws.cell(row=row, column=7).fill = BAD
                ws.cell(row=row, column=7).font = Font(name=FONT, size=10, bold=True)
            if i == stop:
                cell = ws.cell(row=row, column=8)
                cell.fill = BAD
                cell.font = Font(name=FONT, size=10, bold=True, color="9C0006")
            elif stop and i > stop:
                ws.cell(row=row, column=8).font = Font(name=FONT, size=10, color="808080")
            for c in (9, 10):
                ws.cell(row=row, column=c).fill = FILLIN
            row += 1

    # ------------------------------------------------------------------ Scenes
    ws = wb.create_sheet("Scenes")
    head = ["씬", "줄기", "토마토", "높이 (m)", "최소치수 (m)", "익음", "오라클 판정", "기대 동작"]
    ws.append(head)
    style_header(ws, 1, len(head))
    widths(ws, {"A": 6, "B": 10, "C": 12, "D": 10, "E": 13, "F": 9, "G": 14, "H": 14})
    row = 2
    for seed in sorted(scenes):
        for t in scenes[seed]:
            if not t["graspable"]:
                verdict = "불가능 (도달)" if t["violates"] == "too_high" else "불가능 (크기)"
                action = "건너뛰기"
            else:
                verdict = "가능"
                action = "수집" if t["ripeness"] == "ripe" else "방치"
            ws.append([seed, t["stem"][:7], t["name"], t["z"], t["w"], t["ripeness"], verdict, action])
            for c in range(1, len(head) + 1):
                cell = ws.cell(row=row, column=c)
                cell.font = Font(name=FONT, size=10)
                cell.border = BOX
                cell.alignment = Alignment(horizontal="center")
                if seed % 2 == 0:
                    cell.fill = BAND
            if not t["graspable"]:
                ws.cell(row=row, column=7).fill = BAD
            row += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT.relative_to(PROJECT_ROOT)}")
    for name in wb.sheetnames:
        print(f"  {name:16} {wb[name].max_row - 1} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
